import io
import openpyxl
from typing import Dict, Any, List
from app.services.parsers.base import BaseDocumentParser

class XlsxParser(BaseDocumentParser):
    def parse(self, file_bytes: bytes, filename: str) -> Dict[str, Any]:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sections = []
            full_text_parts = []

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_lines = [f"# ÇALIŞMA SAYFASI: {sheet_name}"]
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue

                # Header row
                header = [str(c or "").strip() for c in rows[0]]
                sheet_lines.append("Başlıklar: " + " | ".join(header))

                # Data rows
                for row_idx, row in enumerate(rows[1:], start=2):
                    if any(c is not None for c in row):
                        row_vals = [str(c or "").strip() for c in row]
                        sheet_lines.append(f"Satır {row_idx}: " + " | ".join(row_vals))

                sheet_text = "\n".join(sheet_lines)
                sections.append({"sheet": sheet_name, "content": sheet_text})
                full_text_parts.append(sheet_text)

            return {
                "full_text": "\n\n".join(full_text_parts),
                "sections": sections,
                "metadata": {"format": "xlsx", "sheets": wb.sheetnames}
            }
        except Exception as e:
            raw = file_bytes.decode("utf-8", errors="ignore")
            return {
                "full_text": raw,
                "sections": [{"content": raw}],
                "metadata": {"format": "xlsx_fallback", "error": str(e)}
            }
