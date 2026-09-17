import io
from typing import Dict, Any, List
from app.services.parsers.base import BaseDocumentParser

class DocxParser(BaseDocumentParser):
    def parse(self, file_bytes: bytes, filename: str) -> Dict[str, Any]:
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            sections = []
            full_text_parts = []

            # Extract paragraphs
            para_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            if para_texts:
                full_text_parts.extend(para_texts)
                sections.append({"type": "paragraphs", "content": "\n".join(para_texts)})

            # Extract tables
            for t_idx, table in enumerate(doc.tables, start=1):
                table_lines = []
                for row in table.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_lines.append(" | ".join(row_cells))
                table_text = f"--- Tablo {t_idx} ---\n" + "\n".join(table_lines)
                full_text_parts.append(table_text)
                sections.append({"type": "table", "index": t_idx, "content": table_text})

            return {
                "full_text": "\n\n".join(full_text_parts),
                "sections": sections,
                "metadata": {"format": "docx", "tables_count": len(doc.tables)}
            }
        except Exception:
            # Fallback if docx structure error
            raw = file_bytes.decode("utf-8", errors="ignore")
            return {
                "full_text": raw,
                "sections": [{"content": raw}],
                "metadata": {"format": "docx_fallback"}
            }
