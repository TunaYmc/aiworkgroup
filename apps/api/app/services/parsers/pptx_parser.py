import io
from typing import Dict, Any, List
from app.services.parsers.base import BaseDocumentParser

class PptxParser(BaseDocumentParser):
    def parse(self, file_bytes: bytes, filename: str) -> Dict[str, Any]:
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_bytes))
            sections = []
            full_text_parts = []

            for idx, slide in enumerate(prs.slides, start=1):
                slide_lines = [f"# SLAYT {idx}"]
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_lines.append(shape.text.strip())

                # Check notes
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_lines.append(f"Notlar: {notes}")

                slide_text = "\n".join(slide_lines)
                sections.append({"slide": idx, "content": slide_text})
                full_text_parts.append(slide_text)

            return {
                "full_text": "\n\n".join(full_text_parts),
                "sections": sections,
                "metadata": {"format": "pptx", "slides_count": len(prs.slides)}
            }
        except Exception:
            # Fallback
            raw = file_bytes.decode("utf-8", errors="ignore")
            return {
                "full_text": raw,
                "sections": [{"content": raw}],
                "metadata": {"format": "pptx_fallback"}
            }
